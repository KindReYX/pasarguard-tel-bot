from __future__ import annotations

import tempfile
import asyncio
from pathlib import Path

from aiogram import Router, F
from aiogram.types import FSInputFile, InlineKeyboardButton, InlineKeyboardMarkup, Message, CallbackQuery
from openpyxl import Workbook

import db
from keyboards import EditableButtonFilter
from keyboards import back_keyboard, BTN_EXPORTS
from permissions import is_admin, deny_message, deny_callback


exports_router = Router()


def exports_keyboard() -> InlineKeyboardMarkup:
    return InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text='👥 خروجی کاربران', callback_data='export:users')],
        [InlineKeyboardButton(text='🧾 خروجی سفارش‌ها', callback_data='export:orders')],
    ])


def make_xlsx(rows: list[dict], filename: str) -> Path:
    out = Path(filename)
    wb = Workbook(write_only=True)
    ws = wb.create_sheet()
    if rows:
        headers = list(rows[0].keys())
        ws.append(headers)
        for row in rows:
            ws.append([row.get(h) for h in headers])
    else:
        ws.append(['empty'])
    wb.save(out)
    return out


async def _export_table(callback: CallbackQuery, table: str, order_by: str, caption: str) -> None:
    # Use streaming with limit to avoid OOM; openpyxl write_only is memory-efficient
    # Fetch in chunks of 1000
    offset = 0
    chunk = 1000
    all_rows: list[dict] = []
    # But for simplicity and to avoid loading all into memory twice, we stream directly to workbook
    # Here we create Workbook and append chunk-by-chunk
    tmp = tempfile.NamedTemporaryFile(delete=False, suffix='.xlsx')
    tmp_path = Path(tmp.name)
    tmp.close()
    wb = Workbook(write_only=True)
    ws = wb.create_sheet()
    headers_written = False
    total = 0
    while True:
        rows = db.fetchall(f'SELECT * FROM {table} ORDER BY {order_by} DESC LIMIT ? OFFSET ?', (chunk, offset))
        if not rows and total == 0:
            ws.append(['empty'])
            break
        if rows and not headers_written:
            ws.append(list(rows[0].keys()))
            headers_written = True
        for row in rows:
            ws.append([row.get(h) for h in (list(row.keys()) if rows else [])])
            total += 1
        if len(rows) < chunk:
            break
        offset += chunk
        # Yield to event loop every chunk
        await asyncio.sleep(0)
    wb.save(tmp_path)
    try:
        await callback.message.answer_document(FSInputFile(tmp_path, filename=f'{table}_export.xlsx'), caption=f'{caption} ({total} ردیف)')
    finally:
        try:
            tmp_path.unlink(missing_ok=True)
        except Exception:
            pass


@exports_router.message(EditableButtonFilter(BTN_EXPORTS))
async def exports_home(message: Message):
    if not is_admin(message.from_user.id if message.from_user else None):
        await deny_message(message)
        return
    await message.answer('📤 خروجی اکسل', reply_markup=back_keyboard())
    await message.answer('نوع خروجی را انتخاب کنید:', reply_markup=exports_keyboard())


@exports_router.callback_query(F.data == 'export:users')
async def export_users(callback: CallbackQuery):
    if not is_admin(callback.from_user.id):
        await deny_callback(callback)
        return
    await callback.answer()
    await _export_table(callback, 'bot_users', 'created_at', 'خروجی کاربران')


@exports_router.callback_query(F.data == 'export:orders')
async def export_orders(callback: CallbackQuery):
    if not is_admin(callback.from_user.id):
        await deny_callback(callback)
        return
    await callback.answer()
    await _export_table(callback, 'orders', 'created_at', 'خروجی سفارش‌ها')
